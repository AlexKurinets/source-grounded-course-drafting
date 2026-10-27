from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from pptx import Presentation

from coursedraft.export import ExportError, NotApproved, export_approved, write_deck
from coursedraft.models import Draft, Slide
from coursedraft.synthetic import LEASES
from coursedraft.text import shingles
from coursedraft.validate import facilitator_shingles

COURSE = LEASES["id"]


def exported(rt) -> dict:
    run = rt.start_run(COURSE)
    run = rt.decide(run["id"], "approved", "Dana Whitfield")
    return run


def test_export_refuses_a_draft_without_an_approval_record(leases):
    run = leases.start_run(COURSE)
    with pytest.raises(NotApproved):
        export_approved(leases, run["draft_version_id"])
    returned = leases.decide(run["id"], "returned", "Dana", "Not yet.")
    with pytest.raises(NotApproved):
        export_approved(leases, returned["draft_version_id"])


def test_deck_uses_the_named_layouts_and_keeps_facilitator_text_in_the_notes(leases):
    run = exported(leases)
    deck = next(e for e in run["exports"] if e["kind"] == "deck")
    prs = Presentation(deck["path"])
    drafted = run["draft"]["slides"]
    assert [s.slide_layout.name for s in prs.slides] == [s["layout"] for s in drafted]
    fac = facilitator_shingles(leases.pool, COURSE)
    for s in prs.slides:
        on_slide = " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
        assert not shingles(on_slide) & fac
    notes = prs.slides[3].notes_slide.notes_text_frame.text  # LO2 concept slide
    assert "Sources:" in notes and "standards-update.docx v1" in notes
    footer = [
        sh.text_frame.text
        for sh in prs.slides[3].shapes
        if sh.has_text_frame and sh.text_frame.text.startswith("Source: ")
    ]
    assert footer == ["Source: NASB 116.42; NASB 116.27"]
    assert prs.core_properties.comments.startswith("Approved by Dana Whitfield")


def test_guide_and_case_documents_split_instructor_and_participant_material(leases):
    run = exported(leases)
    paths = {e["kind"]: Path(e["path"]) for e in run["exports"]}
    guide = "\n".join(p.text for p in Document(str(paths["guide"])).paragraphs)
    case = "\n".join(p.text for p in Document(str(paths["case"])).paragraphs)
    assert "Instructor-only. Draft version 1, approved by Dana Whitfield" in guide
    assert "Likely misconceptions" in guide and "Facilitator note" in guide
    assert "Figures to work from" in guide
    assert "Figures to work from" not in case and "Decision point" in case
    assert "Quillmoor Freight Services is a regional freight and warehousing business." in case


def test_a_layout_missing_from_the_template_stops_the_export(leases, tmp_path):
    draft = Draft(slides=[Slide(key="x", layout="Wide Panorama", title="Nope")])
    template = leases.settings.template_dir / "course-master.pptx"
    with pytest.raises(ExportError, match="Wide Panorama"):
        write_deck(
            template,
            {"title": "T"},
            draft,
            {},
            {"reviewer": "x", "decided_at": None},
            tmp_path / "out.pptx",
        )
