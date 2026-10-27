"""Synthetic source material for three continuing-education courses.

Everything here is invented: the Northbridge Accounting Standards Board (NASB) and its paragraph
numbers, the companies in the cases, and every figure. The seed script writes these out as real
.pptx and .docx files and ingests them the same way an instructor's files would be.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from pptx import Presentation
from pptx.util import Inches

LEASES = {
    "id": "lease-accounting-update",
    "title": "Lease Accounting Update",
    "duration_minutes": 180,
    "objectives": [
        {
            "id": "LO1",
            "short": "Reassessing the lease term",
            "text": "Apply the revised lease term reassessment rules to renewal and termination "
            "options.",
        },
        {
            "id": "LO2",
            "short": "Remeasuring index-linked payments",
            "text": "Remeasure the lease liability when index-linked lease payments change.",
        },
        {
            "id": "LO3",
            "short": "Short-term and low-value exemptions",
            "text": "Apply and disclose the short-term and low-value lease exemptions.",
        },
    ],
    "publisher_deck": [
        (
            "Leases: the right-of-use model",
            [
                "Lessees recognise a right-of-use asset and a lease liability for most leases.",
                "The liability is the present value of lease payments not yet paid.",
            ],
        ),
        (
            "Determining the lease term",
            [
                "Start with the non-cancellable period of the lease.",
                "Add option periods the lessee is reasonably certain to exercise.",
            ],
        ),
        (
            "Variable lease payments",
            [
                "Payments linked to an index or rate are included in the lease liability.",
                "Other variable payments are expensed as they are incurred.",
            ],
        ),
        (
            "Remeasuring the lease liability",
            [
                "Remeasure for a change in lease term or a change in index-linked payments.",
                "The adjustment goes against the right-of-use asset.",
            ],
        ),
        (
            "Recognition exemptions",
            [
                "Short-term leases of 12 months or less can be kept off the balance sheet.",
                "Leases of low-value assets can be expensed lease by lease.",
            ],
        ),
    ],
    "publisher_table": (
        "Disclosure checklist",
        [
            ["Item", "Disclosure"],
            ["Short-term leases", "Expense for the period"],
            ["Low-value leases", "Expense for the period"],
            ["Variable payments", "Expense not in the liability"],
        ],
    ),
    "own_deck": [
        (
            "When does the lease term change?",
            [
                "Reassess only for significant events within the lessee's control.",
                "Leasehold improvements and subleases are the usual triggers.",
                "A rent review or a landlord's offer is not a trigger on its own.",
            ],
            "Participants often assume any change in market rents forces a reassessment, and it "
            "does not. Expect someone to ask about a landlord's renewal offer. Point back to control: "
            "the offer is outside the lessee's control.",
        ),
        (
            "Discuss: the warehouse fit-out",
            [
                "A lessee installs racking in year 3 of a 5-year lease with a 5-year renewal option.",
                "Does the fit-out change the lease term, and from which date?",
            ],
            "Expect most groups to say yes. The stronger answer turns on whether the racking only "
            "has value if the lessee stays. A common mistake is to reassess at the board approval "
            "date rather than when the racking is installed.",
        ),
        (
            "Index-linked payments in practice",
            [
                "Measure index-linked payments using the index at commencement.",
                "Remeasure when the adjusted payment takes effect, not when the index is published.",
                "Keep the original discount rate.",
            ],
            "Participants often confuse the publication of a new index figure with the date the "
            "payments change. Walk through a CPI uplift that applies from the next rent date.",
        ),
        (
            "Discuss: the CPI uplift",
            [
                "Rent rises by the change in CPI every two years.",
                "The index rose last month, but the next rent date is in April.",
                "When do you remeasure, and against which account?",
            ],
            "Expect answers that remeasure immediately. The anticipated answer is to wait for the "
            "April rent date and adjust the liability and the right-of-use asset together.",
        ),
        (
            "Using the exemptions well",
            [
                "Short-term means 12 months or less with no purchase option, elected by asset class.",
                "Low-value is assessed on the asset when new, lease by lease.",
            ],
            "A frequent mistake is treating a 12-month lease with a likely renewal as short-term. "
            "Ask the group what happens when the lease is extended.",
        ),
        (
            "Discuss: laptops and vans",
            [
                "A firm leases 400 laptops and 12 delivery vans on 11-month terms.",
                "Which exemptions apply, and what has to be disclosed?",
            ],
            "Expect agreement on the laptops. The vans depend on whether renewal is reasonably "
            "certain; if it is, the term is longer than 12 months.",
        ),
    ],
    "standards": [
        (
            "NASB 116.18 Lease term at commencement",
            [
                "NASB 116.18 A lessee determines the lease term as the non-cancellable period plus "
                "periods covered by an option to extend that it is reasonably certain to exercise. "
                "Periods covered by an option to terminate are included when the lessee is reasonably "
                "certain not to exercise that option."
            ],
        ),
        (
            "NASB 116.20 Reassessment of the lease term",
            [
                "NASB 116.20 A lessee reassesses the lease term only on a significant event or change "
                "in circumstances that is within its control and affects whether it is reasonably "
                "certain to exercise a renewal option. Market-wide changes in rents do not by "
                "themselves trigger a reassessment."
            ],
        ),
        (
            "NASB 116.21 Examples of triggering events",
            [
                "NASB 116.21 Triggering events include significant leasehold improvements not "
                "anticipated at commencement, a significant modification of the leased asset, and a "
                "decision to sublease the asset beyond the current term."
            ],
        ),
        (
            "NASB 116.27 Index-linked payments",
            [
                "NASB 116.27 Lease payments that depend on an index or rate are measured using the "
                "index or rate at the commencement date. Future changes in the index are not "
                "anticipated in the initial measurement."
            ],
        ),
        (
            "NASB 116.42 Remeasurement for a change in payments",
            [
                "NASB 116.42 A lessee remeasures the lease liability when cash flows change because "
                "of a change in an index or rate. The remeasured liability uses the revised payments "
                "and the unchanged discount rate, with the adjustment recorded against the "
                "right-of-use asset."
            ],
        ),
        (
            "NASB 116.43 Timing of remeasurement",
            [
                "NASB 116.43 The remeasurement is made only when the change in cash flows takes "
                "effect, which is when the adjusted payments begin. A published index change that has "
                "not yet altered a payment does not trigger remeasurement."
            ],
        ),
        (
            "NASB 116.60 Short-term lease exemption",
            [
                "NASB 116.60 A lessee may elect not to recognise a right-of-use asset and lease "
                "liability for a lease of 12 months or less at commencement with no purchase option. "
                "The election is made by class of underlying asset."
            ],
        ),
        (
            "NASB 116.61 Low-value asset exemption",
            [
                "NASB 116.61 A lessee may apply the low-value exemption lease by lease when the "
                "underlying asset, assessed when new, is of low value, such as a laptop or office "
                "furniture. The assessment does not depend on the size of the lessee."
            ],
        ),
        (
            "NASB 116.95 Disclosure of exemptions",
            [
                "NASB 116.95 A lessee that applies either exemption discloses that fact and the "
                "expense recognised in the period for short-term leases and for leases of low-value "
                "assets."
            ],
        ),
    ],
    "standards_table": (
        "Transition summary",
        [
            ["Paragraph", "Effective", "Transition"],
            ["NASB 116.20", "Annual periods from 1 January 2026", "Prospective"],
            ["NASB 116.42", "Annual periods from 1 January 2026", "Prospective"],
        ],
    ),
    # The second version of the standards update adds a paragraph; the seed re-ingests it so the
    # demo shows a source version change between two runs of the same course.
    "standards_v2_extra": [
        (
            "NASB 116.44 Discount rate on remeasurement",
            [
                "NASB 116.44 When the change in payments results from a change in floating interest "
                "rates, the lessee remeasures using a revised discount rate that reflects the new "
                "interest rate."
            ],
        ),
    ],
    "client_notes": [
        (
            "Engagement example: distribution client",
            [
                "In a distribution engagement the client had treated a racking fit-out as routine "
                "maintenance and missed the reassessment. Correcting it moved the lease term from 5 to "
                "10 years. Participants often ask whether the auditor would have caught it; expect the "
                "discussion to turn to documentation of the option assessment."
            ],
        ),
        (
            "Engagement example: CPI timing",
            [
                "A client remeasured every quarter when the index was published, which created swings "
                "that reversed at the next rent date. Moving remeasurement to the rent date removed "
                "four adjusting entries a year."
            ],
        ),
    ],
    "fact_record": {
        "company": "Quillmoor Freight Services",
        "industry": "regional freight and warehousing business",
        "decision_maker": "the finance director",
        "decision_point": "Whether to reassess the lease term on the Northgate warehouse after "
        "the racking fit-out, and what to tell the audit committee.",
        "timeline": [
            "Year 1: 5-year lease signed with a 5-year renewal option",
            "Year 3: board approves the racking fit-out",
            "Year 3, fourth quarter: racking installed",
        ],
        "figures": {
            "annual rent": "$420,000",
            "fit-out cost": "$1.1 million",
            "remaining non-cancellable term": "2 years",
            "incremental borrowing rate": "6%",
        },
        "facts": [
            "The Northgate warehouse lease runs 5 years with a 5-year renewal option at "
            "market rent.",
            "The racking cost $1.1 million and has a useful life of 15 years.",
            "Annual rent is $420,000 and the incremental borrowing rate is 6%.",
        ],
    },
}

REVENUE = {
    "id": "revenue-recognition-refresher",
    "title": "Revenue Recognition Refresher",
    "duration_minutes": 120,
    "objectives": [
        {
            "id": "LO1",
            "short": "Identifying performance obligations",
            "text": "Identify distinct performance obligations in bundled software contracts.",
        },
        {
            "id": "LO2",
            "short": "Variable consideration and the constraint",
            "text": "Estimate variable consideration and apply the constraint.",
        },
        {
            "id": "LO3",
            "short": "Allocating the transaction price",
            "text": "Allocate the transaction price on relative standalone selling prices.",
        },
    ],
    "publisher_deck": [
        (
            "The five-step model",
            [
                "Identify the contract, then the performance obligations.",
                "Determine and allocate the transaction price, then recognise revenue.",
            ],
        ),
        (
            "Performance obligations",
            [
                "A promise is a performance obligation if the good or service is distinct.",
                "A series of substantially similar services can be one obligation.",
            ],
        ),
        (
            "Variable consideration",
            [
                "Rebates, credits and usage-based fees make the price variable.",
                "Estimate using the expected value or the most likely amount.",
            ],
        ),
        (
            "Allocating the price",
            [
                "Allocate on relative standalone selling prices.",
                "Estimate prices that are not observable.",
            ],
        ),
    ],
    "publisher_table": (
        "Estimation methods",
        [
            ["Method", "Use when"],
            ["Expected value", "Many similar contracts"],
            ["Most likely amount", "Two possible outcomes"],
        ],
    ),
    "own_deck": [
        (
            "Is the implementation distinct?",
            [
                "Ask whether the customer can benefit from the service on its own or with resources "
                "it has.",
                "Ask whether the promise is separately identifiable in the contract.",
                "Significant customisation usually means it is not distinct.",
            ],
            "Participants often assume that a separately priced service is distinct. Expect the "
            "group to anchor on the invoice; steer them to the integration test instead.",
        ),
        (
            "Discuss: the bundled licence",
            [
                "A software firm sells a licence, implementation and two years of support.",
                "How many performance obligations are there?",
            ],
            "Expect three as the first answer. The anticipated better answer depends on how much "
            "the implementation modifies the software.",
        ),
        (
            "Applying the constraint",
            [
                "Include variable consideration only to the extent a significant reversal is highly "
                "unlikely.",
                "Reassess the estimate at each reporting date.",
            ],
            "A common mistake is to treat the constraint as a reason to exclude all variable "
            "consideration. Walk through a rebate with a long history of similar contracts.",
        ),
        (
            "Discuss: the usage rebate",
            [
                "A customer earns a rebate of up to 10% if usage passes a threshold.",
                "How much of the rebate do you estimate at contract inception?",
            ],
            "Expect a split between expected value and most likely amount. Both can be defended; "
            "push the group to say which fits the data they have.",
        ),
        (
            "Standalone selling prices",
            [
                "Use observable prices where they exist.",
                "Otherwise use adjusted market assessment, expected cost plus margin, or the residual "
                "approach in narrow cases.",
            ],
            "Participants often reach for the residual approach first. It is only available when "
            "prices are highly variable or uncertain.",
        ),
    ],
    "standards": [
        (
            "NASB 115.22 Identifying performance obligations",
            [
                "NASB 115.22 At contract inception an entity identifies as a performance obligation "
                "each promise to transfer a distinct good or service, or a series of substantially the "
                "same distinct goods or services."
            ],
        ),
        (
            "NASB 115.27 Distinct goods or services",
            [
                "NASB 115.27 A good or service is distinct if the customer can benefit from it on its "
                "own or with readily available resources, and the promise is separately identifiable "
                "from other promises in the contract."
            ],
        ),
        (
            "NASB 115.53 Estimating variable consideration",
            [
                "NASB 115.53 An entity estimates variable consideration using the expected value or "
                "the most likely amount, whichever better predicts the consideration it will be "
                "entitled to."
            ],
        ),
        (
            "NASB 115.56 The constraint",
            [
                "NASB 115.56 Variable consideration is included in the transaction price only to the "
                "extent that it is highly probable a significant reversal of cumulative revenue will "
                "not occur when the uncertainty is resolved."
            ],
        ),
        (
            "NASB 115.74 Allocation on relative standalone selling prices",
            [
                "NASB 115.74 The transaction price is allocated to each performance obligation in "
                "proportion to its standalone selling price at contract inception."
            ],
        ),
        (
            "NASB 115.79 Estimating standalone selling prices",
            [
                "NASB 115.79 Where a standalone selling price is not observable, an entity estimates "
                "it using an adjusted market assessment, expected cost plus a margin, or, in limited "
                "circumstances, a residual approach."
            ],
        ),
    ],
    "standards_table": (
        "Effective dates",
        [
            ["Paragraph", "Change", "Effective"],
            ["NASB 115.56", "Wording of the constraint clarified", "1 January 2026"],
            ["NASB 115.79", "Residual approach narrowed", "1 January 2026"],
        ],
    ),
    "standards_v2_extra": [],
    "client_notes": [
        (
            "Engagement example: software reseller",
            [
                "A reseller had recognised implementation fees upfront as a separate obligation even "
                "though the work customised the core product. Combining the obligations deferred "
                "revenue across the implementation period. Expect questions about how the auditor "
                "tested the customisation judgement."
            ],
        ),
    ],
    "fact_record": {
        "company": "Vantorra Software Group",
        "industry": "mid-sized provider of warehouse management software",
        "decision_maker": "the group financial controller",
        "decision_point": "How many performance obligations the Atlas contract contains and how "
        "to allocate its price before the half-year close.",
        "timeline": [
            "January: Atlas contract signed",
            "February to May: implementation and configuration",
            "June: go-live and start of support",
        ],
        "figures": {
            "contract price": "$900,000",
            "licence list price": "$600,000",
            "implementation fee": "$250,000",
            "annual support fee": "$150,000",
            "maximum usage rebate": "10%",
        },
        "facts": [
            "The Atlas contract bundles a licence, implementation and two years of support "
            "for $900,000.",
            "The licence list price is $600,000 and the implementation fee is $250,000.",
            "Support is sold separately for $150,000 a year, and a usage rebate of up to "
            "10% applies.",
        ],
    },
}

IMPAIRMENT = {
    "id": "impairment-essentials",
    "title": "Impairment of Non-Financial Assets",
    "duration_minutes": 120,
    "objectives": [
        {
            "id": "LO1",
            "short": "Spotting impairment indicators",
            "text": "Identify internal and external impairment indicators for assets.",
        },
        {
            "id": "LO2",
            "short": "Measuring recoverable amount",
            "text": "Measure recoverable amount using value in use and fair value less costs of "
            "disposal.",
        },
        {
            "id": "LO3",
            "short": "Recognising and reversing losses",
            "text": "Recognise impairment losses and reversals, including for goodwill.",
        },
    ],
    "publisher_deck": [
        (
            "Why impairment testing exists",
            [
                "Assets should not be carried above the amount they will recover.",
                "Testing is triggered by indicators, and annually for goodwill.",
            ],
        ),
        (
            "Recoverable amount",
            [
                "Recoverable amount is the higher of value in use and fair value less costs of "
                "disposal.",
                "Only one of the two needs to exceed carrying amount.",
            ],
        ),
        (
            "Cash-generating units",
            [
                "Test the smallest group of assets with largely independent cash inflows.",
                "Goodwill is allocated to the units that benefit from it.",
            ],
        ),
        (
            "Reversals",
            [
                "Losses on assets other than goodwill can reverse.",
                "A reversal cannot exceed the carrying amount without the original loss.",
            ],
        ),
    ],
    "publisher_table": (
        "Indicator examples",
        [
            ["External", "Internal"],
            ["Market interest rates rise", "Physical damage"],
            ["Adverse market change", "Plans to restructure"],
        ],
    ),
    "own_deck": [
        (
            "Reading the indicators",
            [
                "Look outside first: interest rates, market values and the competitive landscape.",
                "Then inside: damage, idle capacity and restructuring plans.",
            ],
            "Participants often assume a market capitalisation below net assets is always an "
            "indicator. It is one, but expect someone to argue it away; ask what evidence would.",
        ),
        (
            "Discuss: the bottling line",
            [
                "A bottling plant runs at 60% capacity after losing its largest customer.",
                "Is that an indicator, and for which unit?",
            ],
            "Expect most groups to say yes. The anticipated debate is about the unit: the plant "
            "alone or the plant with the distribution network.",
        ),
        (
            "Value in use",
            [
                "Project cash flows for at most five years before a terminal value.",
                "Use a pre-tax rate that reflects the risks of the asset.",
            ],
            "A common mistake is to include cash flows from improvements the entity has not yet "
            "committed to. Point to the current-condition rule.",
        ),
        (
            "Discuss: two measures",
            [
                "Value in use comes out below carrying amount, but fair value less costs of disposal "
                "is unknown.",
                "Do you need to estimate it before booking a loss?",
            ],
            "Expect the group to book the loss immediately. The anticipated answer is that a fair "
            "value estimate could still show no impairment.",
        ),
        (
            "Allocating a unit's loss",
            [
                "Reduce goodwill first, then the other assets pro rata.",
                "Do not reduce an asset below its own recoverable amount.",
            ],
            "Participants often confuse the order of allocation. Draw the waterfall on the board "
            "before showing the slide.",
        ),
    ],
    "standards": [
        (
            "NASB 136.9 Assessing for indicators",
            [
                "NASB 136.9 At each reporting date an entity assesses whether there is any indication "
                "that an asset may be impaired, and if so estimates its recoverable amount."
            ],
        ),
        (
            "NASB 136.12 Sources of information",
            [
                "NASB 136.12 Indications include a significant decline in market value, adverse "
                "changes in the market or economic environment, increased market interest rates, and "
                "evidence of obsolescence or physical damage."
            ],
        ),
        (
            "NASB 136.18 Recoverable amount",
            [
                "NASB 136.18 Recoverable amount is the higher of an asset's fair value less costs of "
                "disposal and its value in use."
            ],
        ),
        (
            "NASB 136.33 Cash flow projections",
            [
                "NASB 136.33 Projections cover at most five years unless a longer period can be "
                "justified, and reflect the asset in its current condition."
            ],
        ),
        (
            "NASB 136.104 Allocating an impairment loss",
            [
                "NASB 136.104 An impairment loss for a cash-generating unit reduces goodwill first and "
                "then the other assets of the unit pro rata on their carrying amounts."
            ],
        ),
        (
            "NASB 136.114 Reversing an impairment loss",
            [
                "NASB 136.114 An impairment loss on an asset other than goodwill is reversed if the "
                "estimates used to determine recoverable amount have changed since the loss was "
                "recognised."
            ],
        ),
        (
            "NASB 136.124 No reversal for goodwill",
            [
                "NASB 136.124 An impairment loss recognised for goodwill is not reversed in a later "
                "period."
            ],
        ),
    ],
    "standards_table": (
        "Transition summary",
        [
            ["Paragraph", "Change", "Effective"],
            ["NASB 136.33", "Five-year limit restated", "1 January 2026"],
        ],
    ),
    "standards_v2_extra": [],
    "client_notes": [
        (
            "Engagement example: beverage client",
            [
                "A beverage client tested a single plant on its own when its cash inflows depended on "
                "a shared distribution network. Retesting at the network level removed the loss. "
                "Expect the group to ask how the unit was documented."
            ],
        ),
    ],
    "fact_record": {
        "company": "Feldmere Bottling Co.",
        "industry": "regional soft-drink bottler",
        "decision_maker": "the chief financial officer",
        "decision_point": "Whether to recognise an impairment on the Eastfield bottling line at "
        "year end, and how to allocate any loss.",
        "timeline": [
            "March: largest customer gives notice",
            "June: Eastfield line drops to 60% capacity",
            "December: year-end impairment review",
        ],
        "figures": {
            "carrying amount of the unit": "$8.4 million",
            "value in use": "$7.1 million",
            "fair value less costs of disposal": "$6.8 million",
            "goodwill allocated to the unit": "$0.9 million",
            "pre-tax discount rate": "9%",
        },
        "facts": [
            "The Eastfield unit has a carrying amount of $8.4 million, including $0.9 "
            "million of goodwill.",
            "Value in use is $7.1 million at a pre-tax discount rate of 9%.",
            "Fair value less costs of disposal is $6.8 million.",
        ],
    },
}

COURSES = [LEASES, REVENUE, IMPAIRMENT]


def _course_meta(course: dict) -> dict:
    return {k: course[k] for k in ("id", "title", "duration_minutes", "objectives")}


def write_deck(path: Path, slides: list, table: tuple | None = None) -> None:
    prs = Presentation()
    layout = prs.slide_layouts[1]  # Title and Content
    for item in slides:
        title, bullets = item[0], item[1]
        notes = item[2] if len(item) > 2 else ""
        s = prs.slides.add_slide(layout)
        s.shapes.title.text = title
        tf = s.placeholders[1].text_frame
        tf.text = bullets[0]
        for b in bullets[1:]:
            tf.add_paragraph().text = b
        if notes:
            s.notes_slide.notes_text_frame.text = notes
    if table:
        title, rows = table
        s = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
        s.shapes.title.text = title
        shape = s.shapes.add_table(
            len(rows), len(rows[0]), Inches(0.8), Inches(1.8), Inches(8.4), Inches(0.4) * len(rows)
        )
        for r, row in enumerate(rows):
            for c, text in enumerate(row):
                shape.table.cell(r, c).text = text
    prs.save(str(path))


def write_doc(path: Path, title: str, sections: list, table: tuple | None = None) -> None:
    doc = Document()
    doc.add_heading(title, 0)
    for heading, paras in sections:
        doc.add_heading(heading, 1)
        for p in paras:
            doc.add_paragraph(p)
    if table:
        heading, rows = table
        doc.add_heading(heading, 1)
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))
        for r, row in enumerate(rows):
            for c, text in enumerate(row):
                t.cell(r, c).text = text
    doc.save(str(path))


def write_standards(course: dict, root: Path, version: int = 1) -> Path:
    path = root / course["id"] / f"{course['id']}-standards-update.docx"
    path.parent.mkdir(parents=True, exist_ok=True)
    sections = list(course["standards"])
    if version > 1:
        sections += course["standards_v2_extra"]
    write_doc(path, f"{course['title']}: standards update", sections, course["standards_table"])
    return path


def write_course_files(course: dict, root: Path) -> list[tuple]:
    """Write one course's source files. Returns (path, source_type) pairs in ingest order."""
    d = root / course["id"]
    d.mkdir(parents=True, exist_ok=True)
    cid = course["id"]
    publisher = d / f"{cid}-publisher-slides.pptx"
    own = d / f"{cid}-instructor-deck-last-year.pptx"
    notes = d / f"{cid}-instructor-client-notes.docx"
    write_deck(publisher, course["publisher_deck"], course["publisher_table"])
    write_deck(own, course["own_deck"])
    standards = write_standards(course, root)
    write_doc(notes, "Instructor notes from client engagements", course["client_notes"])
    return [
        (publisher, "publisher_deck"),
        (own, "own_deck"),
        (standards, "standard"),
        (notes, "client_notes"),
    ]
