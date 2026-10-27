"""Parse source files into chunks and record them in the source register.

Chunks follow the document's own boundaries: one chunk per slide (plus one for its speaker
notes), one per heading section in a Word file. Tables stay inside their section. Fixed-size token
windows were tried first in production and cut tables and slides in half.
"""

from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.table import Table
from pptx import Presentation
from psycopg_pool import ConnectionPool

from .clock import Clock
from .db import vector_literal
from .embed import Embedder
from .text import paragraph_refs, sha256_bytes, tokens

SOURCE_TYPES = ("own_deck", "publisher_deck", "standard", "client_notes")
# Sensitivity decides the model route: restricted material only ever goes to the local model.
DEFAULT_SENSITIVITY = {
    "publisher_deck": "public",
    "standard": "public",
    "own_deck": "internal",
    "client_notes": "restricted",
}


class IngestError(Exception):
    pass


@dataclass
class ParsedChunk:
    kind: str
    locator: str
    title: str
    body: str
    audience: str


def parse_pptx(path: Path) -> list[ParsedChunk]:
    prs = Presentation(str(path))
    out: list[ParsedChunk] = []
    for i, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None else ""
        parts: list[str] = []
        for shape in slide.shapes:
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text.strip())
            elif shape.has_table:
                rows = [" | ".join(c.text.strip() for c in row.cells) for row in shape.table.rows]
                parts.append("\n".join(rows))
        body = "\n".join(parts) or title
        if body:
            out.append(ParsedChunk("slide", f"slide {i}", title, body, "participant"))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                out.append(ParsedChunk("notes", f"slide {i} notes", title, notes, "facilitator"))
    return out


def parse_docx(path: Path, audience: str = "participant") -> list[ParsedChunk]:
    doc = Document(str(path))
    out: list[ParsedChunk] = []
    title, buf = "", []

    def flush() -> None:
        if buf:
            n = len(out) + 1
            out.append(ParsedChunk("section", f"section {n}", title, "\n".join(buf), audience))
            buf.clear()

    for block in doc.iter_inner_content():
        if isinstance(block, Table):
            buf.append("\n".join(" | ".join(c.text.strip() for c in r.cells) for r in block.rows))
            continue
        style = block.style.name if block.style is not None else ""
        text = block.text.strip()
        if style.startswith("Heading") or style == "Title":
            flush()
            title = text
        elif text:
            buf.append(text)
    flush()
    return out


def parse_pdf(path: Path, audience: str = "participant") -> list[ParsedChunk]:
    """PDFs go through Docling, which keeps headings and tables. It is an optional extra
    because it pulls in layout models; the demo sources are PowerPoint and Word only."""
    try:
        from docling.document_converter import DocumentConverter
    except ImportError as exc:
        raise IngestError("PDF parsing needs the docling extra: pip install '.[docling]'") from exc
    md = DocumentConverter().convert(str(path)).document.export_to_markdown()
    out: list[ParsedChunk] = []
    title, buf = "", []
    for line in md.splitlines() + ["# "]:
        if line.startswith("#"):
            if buf:
                n = len(out) + 1
                out.append(ParsedChunk("section", f"section {n}", title, "\n".join(buf), audience))
                buf = []
            title = line.lstrip("# ").strip()
        elif line.strip():
            buf.append(line.strip())
    return out


def parse_file(path: Path, source_type: str) -> list[ParsedChunk]:
    audience = "facilitator" if source_type == "client_notes" else "participant"
    suffix = path.suffix.lower()
    if suffix == ".pptx":
        return parse_pptx(path)
    if suffix == ".docx":
        return parse_docx(path, audience)
    if suffix == ".pdf":
        return parse_pdf(path, audience)
    raise IngestError(f"unsupported file type: {path.name}")


def tag_objectives(
    text: str, objectives: list[dict], course_title: str = "", threshold: float = 0.25
) -> list[str]:
    """Tag a chunk with the objectives it is about. Words shared by several objectives or by the
    course title ("lease" in a lease course) say nothing about which objective a chunk serves, so
    only each objective's own words count."""
    words = set(tokens(text))
    counts: dict[str, int] = {}
    for obj in objectives:
        for t in set(tokens(obj["text"])):
            counts[t] = counts.get(t, 0) + 1
    common = {t for t, n in counts.items() if n > 1} | set(tokens(course_title))
    scored = []
    for obj in objectives:
        otoks = set(tokens(obj["text"])) - common
        if otoks:
            score = len(words & otoks) / len(otoks)
            if score >= threshold:
                scored.append((score, obj["id"]))
    return [oid for _, oid in sorted(scored, reverse=True)[:2]]


def upsert_course(pool: ConnectionPool, course: dict) -> None:
    with pool.connection() as conn:
        conn.execute(
            """INSERT INTO courses (id, title, duration_minutes, objectives)
               VALUES (%s, %s, %s, %s)
               ON CONFLICT (id) DO UPDATE SET title = EXCLUDED.title,
                 duration_minutes = EXCLUDED.duration_minutes, objectives = EXCLUDED.objectives""",
            (
                course["id"],
                course["title"],
                course["duration_minutes"],
                json.dumps(course["objectives"]),
            ),
        )


def put_fact_record(pool: ConnectionPool, course_id: str, record: dict) -> None:
    with pool.connection() as conn:
        conn.execute(
            """INSERT INTO fact_records (course_id, record) VALUES (%s, %s)
               ON CONFLICT (course_id) DO UPDATE SET record = EXCLUDED.record""",
            (course_id, json.dumps(record)),
        )


def ingest_file(
    pool: ConnectionPool,
    embedder: Embedder,
    clock: Clock,
    course_id: str,
    path: Path,
    source_type: str,
    sensitivity: str | None = None,
) -> dict:
    """Parse one file into chunks. An unchanged file (same hash) is a no-op; a changed one gets
    the next version number and becomes current, and the earlier version's chunks are kept."""
    if source_type not in SOURCE_TYPES:
        raise IngestError(f"unknown source type {source_type!r}")
    sensitivity = sensitivity or DEFAULT_SENSITIVITY[source_type]
    digest = sha256_bytes(path.read_bytes())
    name = path.name
    with pool.connection() as conn, conn.transaction():
        course = conn.execute(
            "SELECT title, objectives FROM courses WHERE id = %s", (course_id,)
        ).fetchone()
        if course is None:
            raise IngestError(f"unknown course {course_id!r}")
        current = conn.execute(
            "SELECT id, version, sha256 FROM sources WHERE course_id = %s AND path = %s "
            "AND current",
            (course_id, name),
        ).fetchone()
        if current and current["sha256"] == digest:
            return {
                "source_id": current["id"],
                "version": current["version"],
                "unchanged": True,
                "chunks": 0,
            }
        parsed = parse_file(path, source_type)
        version = (current["version"] + 1) if current else 1
        if current:
            conn.execute("UPDATE sources SET current = FALSE WHERE id = %s", (current["id"],))
        source_id = conn.execute(
            """INSERT INTO sources (course_id, path, source_type, sensitivity, sha256, version,
                                    ingested_at)
               VALUES (%s, %s, %s, %s, %s, %s, %s) RETURNING id""",
            (course_id, name, source_type, sensitivity, digest, version, clock.now()),
        ).fetchone()["id"]
        for n, ch in enumerate(parsed, start=1):
            text = f"{ch.title}\n{ch.body}"
            conn.execute(
                """INSERT INTO chunks (id, source_id, course_id, kind, locator, title, body,
                                       audience, objective_ids, paragraph_refs, embedding)
                   VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s::vector)""",
                (
                    f"S{source_id}.{n}",
                    source_id,
                    course_id,
                    ch.kind,
                    ch.locator,
                    ch.title,
                    ch.body,
                    ch.audience,
                    tag_objectives(text, course["objectives"], course["title"]),
                    paragraph_refs(text),
                    vector_literal(embedder.embed(text)),
                ),
            )
    return {"source_id": source_id, "version": version, "unchanged": False, "chunks": len(parsed)}
